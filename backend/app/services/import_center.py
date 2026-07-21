"""Import Center v5.2: единая нормализация и пакетный план импорта.

Общий сервис для старых single-file endpoints и нового batch-импорта —
двух реализаций разбора прайсов в проекте быть не должно.

Ключевые правила (v5.2, исправление опасной семантики v4):
- существующий SKU: пустое поле в файле = НЕ менять значение в БД;
- новый SKU: пустой stock = 0, пустой is_active = true;
- явные stock=0 / is_active=false применяются;
- отсутствующие в файле товары никогда не трогаются;
- никакого LLM: только детерминированные правила и алиасы колонок.
"""
import csv
import io
import json
import re
import unicodedata
import zipfile
from dataclasses import dataclass, field
from pathlib import PurePosixPath

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.config import settings
from app.models.product import Product

# ==================== словари ====================

_STR_FIELDS = ("title", "brand", "category", "subcategory", "condition", "color",
               "memory", "storage", "screen_size", "cpu", "ram", "description")
_BOOL_FIELDS = ("is_hot", "is_available_today", "is_active", "is_new", "on_sale")
_TRUE = {"1", "true", "да", "yes", "y", "истина", "+"}
_FALSE = {"0", "false", "нет", "no", "n", "ложь", "-"}
_CONDITIONS = {"new", "used", "refurbished"}

# Детерминированные алиасы заголовков (рус -> канон). Никакого угадывания LLM.
HEADER_ALIASES = {
    "артикул": "sku", "название": "title", "наименование": "title",
    "цена": "price", "старая цена": "old_price", "остаток": "stock",
    "количество": "stock", "кол-во": "stock", "бренд": "brand",
    "категория": "category", "подкатегория": "subcategory",
    "описание": "description", "состояние": "condition", "цвет": "color",
    "память": "memory", "накопитель": "storage", "гарантия": "warranty_months",
}

_KNOWN_COLUMNS = {"sku", "price", "old_price", "stock", "warranty_months",
                  "specs", "tags", "image_1", "image_2", "image_3",
                  *_STR_FIELDS, *_BOOL_FIELDS}

# Служебные листы XLSX, которые не являются данными
_SKIP_SHEETS = {"readme", "сводка", "контроль", "настройки"}

DATA_EXTENSIONS = (".csv", ".json", ".xlsx")
IMAGE_EXTENSIONS = {".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".png": "image/png",
                    ".webp": "image/webp", ".gif": "image/gif"}

# Магические байты: защита от ложного MIME (файл .jpg с не-картинкой внутри)
_MAGIC = [
    (b"\xff\xd8\xff", "image/jpeg"),
    (b"\x89PNG\r\n\x1a\n", "image/png"),
    (b"GIF87a", "image/gif"), (b"GIF89a", "image/gif"),
]

_URL_RE = re.compile(r"^(https?://|/)[^\s]+$")


# ==================== парсеры значений ====================

def parse_bool(value) -> bool | None:
    """Явное значение -> bool; пусто/непонятно -> None (= «не менять»)."""
    if value is None:
        return None
    if isinstance(value, bool):
        return value
    s = str(value).strip().lower()
    if not s:
        return None
    if s in _TRUE:
        return True
    if s in _FALSE:
        return False
    return None


def parse_price(value) -> float | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).strip().replace("\xa0", "").replace(" ", "").replace("₽", "").replace(",", ".")
    if not s:
        return None
    try:
        return float(s)
    except ValueError:
        return None


def parse_json_field(value, kind: str):
    if value is None or value == "":
        return None
    if isinstance(value, dict) and kind == "dict":
        return value
    if isinstance(value, list) and kind == "list":
        return [str(v) for v in value]
    s = str(value).strip()
    try:
        parsed = json.loads(s)
        if kind == "dict" and isinstance(parsed, dict):
            return parsed
        if kind == "list" and isinstance(parsed, list):
            return [str(v) for v in parsed]
    except (ValueError, TypeError):
        pass
    if kind == "list":
        return [t.strip() for t in s.split(",") if t.strip()]
    return "__invalid__"   # dict не распарсился — предупреждение уровнем выше


# ==================== чтение файлов ====================

def canon_key(key: str) -> str:
    k = unicodedata.normalize("NFC", str(key or "")).strip().lower()
    return HEADER_ALIASES.get(k, k)


def rows_from_csv(data: bytes) -> list[dict]:
    text = data.decode("utf-8-sig", errors="replace")
    sample = text[:2048]
    delimiter = ";" if sample.count(";") > sample.count(",") else ","
    return list(csv.DictReader(io.StringIO(text), delimiter=delimiter))


def rows_from_json(data: bytes) -> list[dict]:
    parsed = json.loads(data.decode("utf-8-sig", errors="replace"))
    items = parsed if isinstance(parsed, list) else parsed.get("items") if isinstance(parsed, dict) else None
    if not isinstance(items, list):
        raise ValueError('Ожидается JSON array товаров или {"items": [...]}')
    return [r for r in items if isinstance(r, dict)]


def rows_from_xlsx(data: bytes) -> tuple[list[dict], str]:
    """Выбор листа (v5.2): лист 'products' -> иначе первый непустой,
    пропуская служебные README/Сводка/Контроль/Настройки."""
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        sheet = None
        for ws in wb.worksheets:
            if ws.title.strip().lower() == "products":
                sheet = ws
                break
        if sheet is None:
            for ws in wb.worksheets:
                if ws.title.strip().lower() in _SKIP_SHEETS:
                    continue
                if ws.max_row and ws.max_row > 1:
                    sheet = ws
                    break
        if sheet is None:
            return [], ""
        rows_iter = sheet.iter_rows(values_only=True)
        header = next(rows_iter, None)
        if not header:
            return [], sheet.title
        keys = [str(h or "").strip().lower() for h in header]
        out = []
        for values in rows_iter:
            if values is None or all(v in (None, "") for v in values):
                continue
            out.append({k: v for k, v in zip(keys, values) if k})
        return out, sheet.title
    finally:
        wb.close()


def read_data_file(filename: str, data: bytes) -> tuple[list[dict], str | None]:
    """-> (rows, sheet_name|None). Бросает ValueError с понятным текстом."""
    name = filename.lower()
    if name.endswith(".csv") or name.endswith(".txt"):
        return rows_from_csv(data), None
    if name.endswith(".json"):
        return rows_from_json(data), None
    if name.endswith(".xlsx"):
        return rows_from_xlsx(data)
    raise ValueError("Поддерживаются файлы: .csv, .json, .xlsx")


# ==================== нормализация строки ====================

def normalize_row(raw: dict) -> tuple[str | None, dict, list[str], list[str], list[str]]:
    """Строка файла -> (sku, changes, errors, warnings, info).

    changes содержит ТОЛЬКО явно заполненные поля (частичная семантика);
    дефолты для новых товаров применяются на этапе плана.
    """
    errors: list[str] = []
    warnings: list[str] = []
    info: list[str] = []
    row = {canon_key(k): v for k, v in raw.items()}

    unknown = [k for k in row if k and k not in _KNOWN_COLUMNS]
    if unknown:
        info.append(f"неизвестные колонки игнорированы: {', '.join(sorted(unknown)[:8])}")
    aliased = [k for k in raw if canon_key(k) != str(k or "").strip().lower() and canon_key(k) in _KNOWN_COLUMNS]
    if aliased:
        info.append("алиасы заголовков: " + ", ".join(f"{k}→{canon_key(k)}" for k in aliased[:6]))

    sku = str(row.get("sku") or "").strip().upper()
    if not sku:
        return None, {}, ["sku обязателен"], [], info
    if len(sku) > 64:
        return None, {}, ["sku длиннее 64 символов"], [], info

    ch: dict = {}
    for f in _STR_FIELDS:
        v = row.get(f)
        if v is not None and str(v).strip() != "":
            ch[f] = str(v).strip()

    if "condition" in ch:
        low = ch["condition"].lower()
        if low not in _CONDITIONS:
            warnings.append(f"condition '{ch['condition']}' не из new/used/refurbished — заменён на new")
            ch["condition"] = "new"
        else:
            ch["condition"] = low

    if row.get("price") not in (None, ""):
        price = parse_price(row.get("price"))
        if price is None:
            errors.append(f"некорректная цена: {row.get('price')!r}")
        elif price <= 0:
            errors.append(f"price должна быть > 0, получено {price}")
        else:
            ch["price"] = price
    if row.get("old_price") not in (None, ""):
        old_price = parse_price(row.get("old_price"))
        if old_price is not None and old_price > 0:
            ch["old_price"] = old_price
            if "price" in ch and old_price <= ch["price"]:
                warnings.append("old_price не больше price — скидка не будет показана")

    if row.get("stock") not in (None, ""):
        try:
            ch["stock"] = max(0, int(float(str(row["stock"]).replace(" ", ""))))
        except (TypeError, ValueError):
            warnings.append(f"некорректный stock {row.get('stock')!r} — поле пропущено")

    if row.get("warranty_months") not in (None, ""):
        try:
            ch["warranty_months"] = max(0, int(float(row["warranty_months"])))
        except (TypeError, ValueError):
            warnings.append("некорректный warranty_months — пропущен")

    for f in _BOOL_FIELDS:
        v = parse_bool(row.get(f))
        if v is not None:
            ch[f] = v

    specs = parse_json_field(row.get("specs"), "dict")
    if specs == "__invalid__":
        warnings.append("specs не является JSON-объектом — пропущены")
    elif specs:
        ch["specs"] = specs
    tags = parse_json_field(row.get("tags"), "list")
    if tags and tags != "__invalid__":
        ch["tags"] = tags

    images = [str(row[k]).strip() for k in ("image_1", "image_2", "image_3")
              if row.get(k) and str(row[k]).strip()]
    if images:
        bad = [u for u in images if not _URL_RE.match(u)]
        if bad:
            warnings.append(f"невалидный image URL: {bad[0][:60]}")
        ch["images"] = [u for u in images if _URL_RE.match(u)]

    return sku, ch, errors, warnings, info


# ==================== безопасный разбор ZIP ====================

@dataclass
class ZipEntry:
    name: str          # базовое имя файла
    path: str          # путь внутри архива (для отчёта)
    data: bytes
    content_type: str


@dataclass
class ZipScan:
    images: list[ZipEntry] = field(default_factory=list)
    data_files: list[tuple[str, bytes]] = field(default_factory=list)  # (имя, данные)
    manifest: dict | None = None
    skipped: list[dict] = field(default_factory=list)   # info: README, скрытые и т.п.
    errors: list[dict] = field(default_factory=list)


def _sniff_image(data: bytes, ext: str) -> str | None:
    """Настоящий тип по магическим байтам; None если содержимое — не картинка."""
    for magic, ctype in _MAGIC:
        if data.startswith(magic):
            return ctype
    if data[:4] == b"RIFF" and data[8:12] == b"WEBP":
        return "image/webp"
    return None


def scan_zip(data: bytes) -> ZipScan:
    """Безопасная распаковка: zip-slip, абсолютные пути, symlink, bomb,
    лимиты количества/объёма, ложный MIME, дубли имён, служебный мусор."""
    scan = ZipScan()
    try:
        zf = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile:
        raise ValueError("Файл не является корректным ZIP-архивом")

    total_uncompressed = 0
    image_count = 0
    seen_names: set[str] = set()

    for entry in zf.infolist():
        if entry.is_dir():
            continue
        raw_path = entry.filename.replace("\\", "/")
        base = PurePosixPath(raw_path).name

        # --- безопасность пути ---
        if raw_path.startswith("/") or re.match(r"^[a-zA-Z]:", raw_path) or ".." in raw_path.split("/"):
            scan.errors.append({"file": raw_path, "error": "недопустимый путь в архиве (zip-slip)"})
            continue
        # symlink: старший байт external_attr — unix mode; S_IFLNK = 0xA000
        if (entry.external_attr >> 16) & 0xF000 == 0xA000:
            scan.errors.append({"file": raw_path, "error": "symlink в архиве запрещён"})
            continue
        # служебный мусор
        if "__MACOSX" in raw_path or base in (".DS_Store", "Thumbs.db") or base.startswith("."):
            scan.skipped.append({"file": raw_path, "info": "служебный файл пропущен"})
            continue

        # --- decompression bomb: заявленный и фактический размер ---
        total_uncompressed += entry.file_size
        if total_uncompressed > settings.IMPORT_MAX_UNCOMPRESSED_BYTES:
            raise ValueError("Суммарный распакованный размер архива превышает лимит")

        ext = ("." + base.rsplit(".", 1)[-1].lower()) if "." in base else ""

        if base.lower() == "manifest.json":
            try:
                scan.manifest = json.loads(zf.read(entry).decode("utf-8-sig"))
            except (ValueError, UnicodeDecodeError):
                scan.errors.append({"file": raw_path, "error": "manifest.json не является валидным JSON"})
            continue

        if ext in DATA_EXTENSIONS:
            if entry.file_size > settings.IMPORT_MAX_DATA_FILE_BYTES:
                scan.errors.append({"file": raw_path, "error": "data-файл больше лимита"})
                continue
            payload = _safe_read(zf, entry, settings.IMPORT_MAX_DATA_FILE_BYTES)
            if payload is None:
                scan.errors.append({"file": raw_path, "error": "фактический размер больше заявленного (bomb)"})
                continue
            scan.data_files.append((base, payload))
            continue

        if ext in IMAGE_EXTENSIONS:
            image_count += 1
            if image_count > settings.IMPORT_MAX_IMAGES:
                raise ValueError(f"Больше {settings.IMPORT_MAX_IMAGES} изображений в архиве")
            if entry.file_size > 8 * 1024 * 1024:
                scan.errors.append({"file": raw_path, "error": "изображение больше 8 МБ"})
                continue
            if base.lower() in seen_names:
                scan.errors.append({"file": raw_path, "error": "повторяющееся имя файла"})
                continue
            seen_names.add(base.lower())
            payload = _safe_read(zf, entry, 8 * 1024 * 1024)
            if payload is None:
                scan.errors.append({"file": raw_path, "error": "фактический размер больше заявленного (bomb)"})
                continue
            ctype = _sniff_image(payload, ext)
            if ctype is None:
                scan.errors.append({"file": raw_path, "error": "содержимое не является изображением (ложный MIME)"})
                continue
            scan.images.append(ZipEntry(name=base, path=raw_path, data=payload, content_type=ctype))
            continue

        scan.skipped.append({"file": raw_path, "info": f"неподдерживаемый файл ({ext or 'без расширения'}) — игнорирован"})

    return scan


def _safe_read(zf: zipfile.ZipFile, entry: zipfile.ZipInfo, limit: int) -> bytes | None:
    """Чтение с жёстким потолком: заголовок ZIP может врать о размере."""
    with zf.open(entry) as fh:
        data = fh.read(limit + 1)
    return None if len(data) > limit else data


# Номер фотографии в галерее. Всё, что больше, — почти наверняка часть SKU
# (год выпуска: -2024, -2025, -2026), а не порядковый номер кадра.
MAX_GALLERY_INDEX = 20


def parse_filename_candidates(name: str) -> list[tuple[str, int]]:
    """Кандидаты (sku, image_index) по убыванию приоритета. Без обращения к БД.

    Порядок принципиален — точный SKU всегда важнее галерейного суффикса:
      1) '<stem>_main' / '<stem>-main'  -> [(stem, 0)] и больше ничего;
      2) весь stem как точный SKU       -> (stem, 0);
      3) хвост '-N' / '_N' при 1<=N<=MAX_GALLERY_INDEX -> (head, N).

    Раньше шаг 3 выполнялся ПЕРВЫМ и без верхней границы, поэтому числовой
    хвост реального SKU принимался за номер кадра: 'APL-APD-4-2024.jpg' читался
    как ('APL-APD-4', 2024), а 'APL-APD-4-2024-2.jpg' — как ('APL-APD-4-2024', 2),
    и снимок уходил чужому товару.
    """
    stem = PurePosixPath(name.replace("\\", "/")).stem
    lower = stem.lower()
    if lower.endswith("_main") or lower.endswith("-main"):
        return [(stem[:-5], 0)]

    out: list[tuple[str, int]] = [(stem, 0)]
    for sep in ("-", "_"):
        head, _, tail = stem.rpartition(sep)
        if head and tail.isdigit():
            n = int(tail)
            if 1 <= n <= MAX_GALLERY_INDEX:
                out.append((head, n))
            break
    return out


def build_sku_index(skus) -> dict[str, str]:
    """{lower(sku): canonical_sku} — строится ОДИН раз на job, не на файл.

    Принимает и {lower: canonical} (значение — каноническое написание), и
    {canonical: что-угодно} (например sku -> id), и просто список SKU.
    Наружу всегда отдаётся каноническое написание из базы, не ключ.
    """
    if isinstance(skus, dict):
        out: dict[str, str] = {}
        for key, val in skus.items():
            canon = val if isinstance(val, str) and val else key
            out[canon.lower()] = canon
        return out
    return {s.lower(): s for s in skus if s}


def resolve_filename_to_sku(name: str, sku_index: dict[str, str]) -> tuple[str | None, int]:
    """Имя файла -> (канонический SKU из каталога | None, image_index).

    sku_index — заранее построенный {lower: canonical}: сравнение
    case-insensitive, наружу отдаётся каноническое написание из базы.
    """
    for guess, order in parse_filename_candidates(name):
        canon = sku_index.get(guess.lower())
        if canon is not None:
            return canon, order
    return None, 0


# ==================== построение пакетного плана ====================

def build_batch_plan(
    db: Session,
    data_files: list[dict],      # [{"name", "rows", "sheet"}]
    image_entries: list[ZipEntry],
    mode: str = "create_or_update",
    duplicate_policy: str = "error",
) -> dict:
    """Merge всех файлов -> нормализованный план + подробный отчёт.

    БД НЕ меняется. Секретные пути наружу не отдаются.
    """
    existing = {
        (p.sku or "").lower(): p
        for p in db.execute(select(Product).where(Product.sku.is_not(None))).scalars()
    }

    rows_out: list[dict] = []
    by_sku: dict[str, dict] = {}
    duplicates: list[dict] = []
    files_report: list[dict] = []

    for f in data_files:
        file_rows = 0
        file_errors = 0
        for i, raw in enumerate(f["rows"]):
            line = i + 2
            sku, changes, errors, warnings, info = normalize_row(raw)
            entry = {
                "source_file": f["name"], "source_sheet": f.get("sheet"),
                "source_line": line, "sku": sku, "action": None,
                "changes": changes, "warnings": warnings, "errors": errors, "info": info,
            }
            file_rows += 1
            if sku is None:
                entry["action"] = "skip"
                file_errors += 1
                rows_out.append(entry)
                continue

            prev = by_sku.get(sku.lower())
            if prev is not None:
                if duplicate_policy == "last_wins":
                    prev["action"] = "superseded"
                    prev["warnings"] = [*prev["warnings"],
                                        f"перекрыт файлом {f['name']} (строка {line}) — last_wins"]
                    duplicates.append({"sku": sku, "loser": prev["source_file"],
                                       "winner": f["name"], "policy": "last_wins"})
                else:
                    entry["errors"] = [*errors,
                                       f"дубль SKU: уже встречался в {prev['source_file']} "
                                       f"(строка {prev['source_line']})"]
                    entry["action"] = "skip"
                    file_errors += 1
                    duplicates.append({"sku": sku, "first": prev["source_file"],
                                       "second": f["name"], "policy": "error"})
                    rows_out.append(entry)
                    continue

            by_sku[sku.lower()] = entry
            rows_out.append(entry)
        files_report.append({"name": f["name"], "sheet": f.get("sheet"),
                             "rows": file_rows, "errors": file_errors})

    # --- определяем действия и диффы ---
    created = updated = unchanged = skipped = 0
    total_price_before = sum(float(p.price) for p in existing.values())
    total_price_delta = 0.0

    for entry in rows_out:
        if entry["action"] in ("skip", "superseded"):
            skipped += 1
            continue
        sku = entry["sku"]
        changes = entry["changes"]
        product = existing.get(sku.lower())

        if product is None:
            if mode == "update_only":
                entry["action"] = "skip"
                entry["warnings"].append("SKU не найден, режим update_only — строка пропущена")
                skipped += 1
                continue
            entry["action"] = "create"
            if not changes.get("title"):
                entry["errors"].append("для нового товара нужен title")
            if changes.get("price") is None:
                entry["errors"].append("для нового товара нужна корректная price")
            # дефолты нового товара (v5.2): пустой stock=0, пустой is_active=true
            changes.setdefault("stock", 0)
            changes.setdefault("is_active", True)
            if changes.get("is_active") and changes.get("stock", 0) == 0:
                entry["warnings"].append("товар активен, но stock=0 — покажется как «под заказ»")
            if entry["errors"]:
                entry["action"] = "skip"
                skipped += 1
                continue
            created += 1
            total_price_delta += changes.get("price") or 0.0
        else:
            if mode == "create_only":
                entry["action"] = "skip"
                entry["warnings"].append("SKU уже существует, режим create_only — строка пропущена")
                skipped += 1
                continue
            # частичная семантика: пустые поля НЕ трогаем; diff только по реальным изменениям
            diff = {}
            for field_name, new_value in changes.items():
                old_value = getattr(product, field_name, None)
                if field_name in ("price", "old_price") and old_value is not None:
                    old_value = float(old_value)
                if field_name == "images":
                    old_value = product.images or []
                if old_value != new_value:
                    diff[field_name] = {"old": old_value, "new": new_value}
            if "price" in diff and diff["price"]["old"]:
                change_pct = abs(diff["price"]["new"] - diff["price"]["old"]) / diff["price"]["old"] * 100
                if change_pct > settings.IMPORT_PRICE_CHANGE_WARN_PCT:
                    entry["warnings"].append(
                        f"цена меняется на {change_pct:.0f}% "
                        f"({diff['price']['old']:.0f} → {diff['price']['new']:.0f})")
                total_price_delta += diff["price"]["new"] - diff["price"]["old"]
            new_stock = changes.get("stock", product.stock)
            new_active = changes.get("is_active", product.is_active)
            if new_active and (new_stock or 0) == 0:
                entry["warnings"].append("товар активен, но stock=0 — покажется как «под заказ»")
            if entry["errors"]:
                entry["action"] = "skip"
                skipped += 1
                continue
            if not diff:
                entry["action"] = "unchanged"
                unchanged += 1
            else:
                entry["action"] = "update"
                entry["diff"] = diff
                updated += 1

    # --- матчинг фото: существующие SKU + SKU, создаваемые этой job ---
    planned_skus = {e["sku"].lower(): e["sku"] for e in rows_out
                    if e["sku"] and e["action"] in ("create", "update", "unchanged")}
    known_skus = {**{k: (existing[k].sku or "") for k in existing}, **planned_skus}

    images_by_sku: dict[str, list[tuple[int, ZipEntry]]] = {}
    unmatched_images: list[str] = []
    sku_index = build_sku_index(known_skus)
    for img in image_entries:
        canon, order = resolve_filename_to_sku(img.name, sku_index)
        if canon is None:
            unmatched_images.append(img.path)
            continue
        images_by_sku.setdefault(canon, []).append((order, img))

    images_report = []
    for sku, files in sorted(images_by_sku.items()):
        files.sort(key=lambda t: t[0])
        images_report.append({
            "sku": sku, "matched_images": len(files),
            "main_file": files[0][1].name,
            "files": [f.name for _, f in files],
        })

    matched_sku_set = {r["sku"].lower() for r in images_report}
    products_without_photos = [
        e["sku"] for e in rows_out
        if e["action"] == "create" and e["sku"] and e["sku"].lower() not in matched_sku_set
        and not e["changes"].get("images")
    ]

    error_rows = [e for e in rows_out if e["errors"]]
    warning_rows = [e for e in rows_out if e["warnings"]]

    catalog_count = len(existing)
    avg_before = (total_price_before / catalog_count) if catalog_count else 0.0
    new_count = catalog_count + created
    avg_after = ((total_price_before + total_price_delta) / new_count) if new_count else 0.0

    return {
        "summary": {
            "files_total": len(data_files) + (1 if image_entries else 0),
            "data_files": len(data_files),
            "zip_images": len(image_entries),
            "rows_total": len(rows_out),
            "create": created, "update": updated, "unchanged": unchanged,
            "skip": skipped, "errors": len(error_rows), "warnings": len(warning_rows),
            "images_matched": sum(r["matched_images"] for r in images_report),
            "images_unmatched": len(unmatched_images),
            "products_without_photos": len(products_without_photos),
            "duplicates": len(duplicates),
            "catalog_total_price_before": round(total_price_before, 2),
            "catalog_total_price_after": round(total_price_before + total_price_delta, 2),
            "catalog_avg_price_before": round(avg_before, 2),
            "catalog_avg_price_after": round(avg_after, 2),
        },
        "files": files_report,
        "rows": rows_out,
        "duplicates": duplicates,
        "images": images_report,
        "unmatched_images": unmatched_images,
        "products_without_photos": products_without_photos[:100],
        "mode": mode,
        "duplicate_policy": duplicate_policy,
    }


# ==================== применение плана ====================

def apply_product_plan(db: Session, plan: dict) -> dict:
    """Применить товарную часть плана ОДНОЙ транзакцией (без commit внутри).

    Возвращает {sku: product_id}. Любая ошибка -> исключение, вызывающий
    делает rollback — частично применённых товаров не бывает.
    """
    existing = {
        (p.sku or "").lower(): p
        for p in db.execute(select(Product).where(Product.sku.is_not(None))).scalars()
    }
    sku_to_id: dict[str, int] = {}
    for entry in plan["rows"]:
        action = entry.get("action")
        if action not in ("create", "update"):
            if action == "unchanged" and entry["sku"]:
                p = existing.get(entry["sku"].lower())
                if p:
                    sku_to_id[entry["sku"]] = p.id
            continue
        changes = dict(entry["changes"])
        images = changes.pop("images", None)
        if action == "create":
            product = Product(title=changes["title"], price=changes["price"], source="import")
            db.add(product)
        else:
            product = existing.get(entry["sku"].lower())
            if product is None:   # товар удалили между preview и confirm
                raise ValueError(f"SKU {entry['sku']} исчез из каталога между preview и confirm")
        for field_name, value in changes.items():
            setattr(product, field_name, value)
        product.sku = entry["sku"]
        if images:
            product.images = images
            product.image = images[0]
        product.in_stock = (product.stock or 0) > 0
        db.flush()
        sku_to_id[entry["sku"]] = product.id
    return sku_to_id


# ==================== режим MODEL_COLOR: галерея на группу модель+цвет ====================

def _basename(name) -> str:
    return str(name or "").replace("\\", "/").split("/")[-1]


def normalize_group_specs(manifest) -> list[dict]:
    """Извлечь описания групп из manifest. Поддержаны две формы:
      {"image_groups": [ {brand, model, color, files, primary}, ... ]}
      {"matchMode": "model_color", brand, model, color, files, primary}   (одна группа)
    """
    if not isinstance(manifest, dict):
        return []
    specs = manifest.get("image_groups")
    if not isinstance(specs, list):
        if manifest.get("matchMode") == "model_color" or ("model" in manifest and "files" in manifest):
            specs = [manifest]
        else:
            return []
    return [s for s in specs if isinstance(s, dict)]


def plan_image_group_import(db: Session, manifest, available_images: set[str]) -> dict:
    """Разобрать manifest MODEL_COLOR -> план по группам. БД НЕ меняется.

    Для каждой группы показываем: распознанные brand/model/color, canonical key,
    какие файлы есть/отсутствуют в архиве, сколько вариантов товара получат эту
    галерею (+ их id/sku), конфликты одинаковых групп, ошибки. При неоднозначном
    или неполном описании группа помечается ошибкой и НЕ применяется молча.
    """
    from app.services.image_groups import image_group_key

    groups_out: list[dict] = []
    seen_keys: dict[str, int] = {}
    for idx, spec in enumerate(normalize_group_specs(manifest)):
        brand = str(spec.get("brand") or "").strip()
        model = str(spec.get("model") or "").strip()
        color = str(spec.get("color") or "").strip()
        files = [_basename(f) for f in (spec.get("files") or [])]
        primary = _basename(spec.get("primary")) if spec.get("primary") else None
        key = image_group_key(brand, model, color=color, model_family=model)
        entry = {
            "index": idx, "brand": brand, "model": model, "color": color, "key": key,
            "files": files, "primary": primary, "present": [], "missing": [],
            "affected_variants": 0, "affected": [], "errors": [], "warnings": [],
        }
        if not key:
            entry["errors"].append("нужны brand, model и color для устойчивого ключа группы")
            groups_out.append(entry)
            continue
        if key in seen_keys:
            entry["errors"].append(f"конфликт: та же группа модель+цвет уже описана (группа #{seen_keys[key] + 1})")
            groups_out.append(entry)
            continue
        seen_keys[key] = idx
        entry["present"] = [f for f in files if f in available_images]
        entry["missing"] = [f for f in files if f not in available_images]
        if not entry["present"]:
            entry["errors"].append("ни одного из указанных файлов нет в архиве")
        if primary and primary not in entry["present"]:
            entry["warnings"].append("primary отсутствует — главной станет первый доступный файл")
        variants = db.execute(select(Product).where(Product.image_group_key == key)).scalars().all()
        entry["affected_variants"] = len(variants)
        entry["affected"] = [{"id": p.id, "sku": p.sku, "title": p.title} for p in variants[:50]]
        if not variants:
            entry["warnings"].append("сейчас нет товаров этой модели+цвета — галерея подключится, когда они появятся")
        groups_out.append(entry)

    ok = [g for g in groups_out if not g["errors"] and g["present"]]
    return {
        "groups": groups_out,
        "summary": {
            "groups_total": len(groups_out),
            "groups_ok": len(ok),
            "groups_with_errors": sum(1 for g in groups_out if g["errors"]),
            "variants_affected": sum(g["affected_variants"] for g in ok),
            "images_total": sum(len(g["present"]) for g in ok),
        },
    }


def apply_image_group_import(db: Session, groups_plan: dict, url_by_filename: dict[str, str]) -> dict:
    """Создать/обновить ProductImageGroup по плану. Порядок: primary первым,
    затем остальные present в порядке files. Существующие URL не удаляем без
    замены — перезаписываем галерею группы целиком новыми файлами."""
    from app.models.product_image_group import ProductImageGroup

    applied: list[dict] = []
    for g in groups_plan["groups"]:
        if g["errors"] or not g["present"]:
            continue
        ordered: list[str] = []
        if g["primary"] and g["primary"] in g["present"]:
            ordered.append(g["primary"])
        for f in g["files"]:
            if f in g["present"] and f not in ordered:
                ordered.append(f)
        urls = [url_by_filename[f] for f in ordered if f in url_by_filename]
        if not urls:
            continue
        grp = db.execute(
            select(ProductImageGroup).where(ProductImageGroup.key == g["key"])
        ).scalar_one_or_none()
        if grp is None:
            grp = ProductImageGroup(key=g["key"], brand=g["brand"], model=g["model"], color=g["color"])
            db.add(grp)
        grp.image = urls[0]
        grp.images = urls
        applied.append({"key": g["key"], "images": len(urls),
                        "affected_variants": g["affected_variants"]})
    db.commit()
    return {"applied_groups": applied, "groups_applied": len(applied)}
