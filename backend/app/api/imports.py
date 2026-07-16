"""Import Center (v4): массовый импорт товаров и фото.

POST /api/admin/import/products/preview — разобрать файл (CSV/JSON/XLSX),
    провалидировать и показать, что будет создано/обновлено. БД не трогает.
POST /api/admin/import/products/confirm — тот же файл, применяет изменения.
POST /api/admin/uploads/products/images-zip — ZIP с фото, авто-матчинг по SKU.

Preview и confirm оба принимают файл (состояние на сервере не храним):
админка сначала шлёт файл на preview, показывает отчёт, по кнопке
«Импортировать» шлёт тот же файл на confirm. Просто и надёжно.

Правила импорта:
- sku обязателен (ключ матчинга);
- sku существует -> обновление, нет -> создание;
- старые товары не удаляются;
- price обязателен и валиден для новых товаров;
- stock пустой -> 0; is_active пустой -> true.
"""
import csv
import io
import json
import zipfile
from pathlib import PurePosixPath

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile, status
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.api.deps import get_current_admin
from app.core.uploads import save_image
from app.db.session import get_db
from app.models.product import Product

router = APIRouter(prefix="/admin", tags=["admin-import"], dependencies=[Depends(get_current_admin)])

MAX_ROWS = 2000
MAX_FILE_BYTES = 20 * 1024 * 1024   # 20 МБ на файл импорта
MAX_ZIP_BYTES = 200 * 1024 * 1024   # 200 МБ на ZIP с фото

# Колонки, которые понимает импорт (лишние молча игнорируем — прайсы бывают разные)
_STR_FIELDS = ("title", "brand", "category", "subcategory", "condition", "color",
               "memory", "storage", "screen_size", "cpu", "ram", "description")
_BOOL_FIELDS = ("is_hot", "is_available_today", "is_active", "is_new", "on_sale")
_TRUE = {"1", "true", "да", "yes", "y", "истина", "+"}
_FALSE = {"0", "false", "нет", "no", "n", "ложь", "-", ""}
_CONDITIONS = {"new", "used", "refurbished"}


def _parse_bool(value, default: bool | None = None) -> bool | None:
    if value is None:
        return default
    if isinstance(value, bool):
        return value
    s = str(value).strip().lower()
    if s in _TRUE:
        return True
    if s in _FALSE:
        return default if s == "" else False
    return default


def _parse_price(value):
    """'89 990,50 ₽' -> 89990.5; пусто/мусор -> None."""
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value) if value > 0 else None
    s = str(value).strip().replace("\xa0", "").replace(" ", "").replace("₽", "").replace(",", ".")
    if not s:
        return None
    try:
        p = float(s)
        return p if p > 0 else None
    except ValueError:
        return None


def _parse_json_field(value, kind: str):
    """specs -> dict, tags -> list. Принимает и JSON-строку, и 'а,б,в' для tags."""
    if value is None or value == "":
        return {} if kind == "dict" else []
    if isinstance(value, dict) and kind == "dict":
        return value
    if isinstance(value, list) and kind == "list":
        return [str(v) for v in value]
    s = str(value).strip()
    try:
        parsed = json.loads(s)
        if kind == "dict" and isinstance(parsed, dict):
            return parsed
        if kind == "list" and isinstance(parsed, list):
            return [str(v) for v in parsed]
    except (ValueError, TypeError):
        pass
    if kind == "list":  # запасной формат: "хит, титан"
        return [t.strip() for t in s.split(",") if t.strip()]
    return None  # dict не распарсился — это warning


def _normalize_row(raw: dict, index: int) -> tuple[dict | None, list[str], list[str]]:
    """Строка файла -> нормализованный payload + errors + warnings."""
    errors: list[str] = []
    warnings: list[str] = []
    row = { (k or "").strip().lower(): v for k, v in raw.items() }

    # SKU канонизируем в верхний регистр: 'iphone15pro' и 'IPHONE15PRO' — один товар
    sku = str(row.get("sku") or "").strip().upper()
    if not sku:
        return None, ["sku обязателен"], []
    if len(sku) > 64:
        return None, ["sku длиннее 64 символов"], []

    item: dict = {"sku": sku}
    for f in _STR_FIELDS:
        v = row.get(f)
        if v is not None and str(v).strip() != "":
            item[f] = str(v).strip()

    if "condition" in item and item["condition"].lower() not in _CONDITIONS:
        warnings.append(f"condition '{item['condition']}' не из new/used/refurbished — заменён на new")
        item["condition"] = "new"
    elif "condition" in item:
        item["condition"] = item["condition"].lower()

    price = _parse_price(row.get("price"))
    if row.get("price") not in (None, "") and price is None:
        errors.append(f"некорректная цена: {row.get('price')!r}")
    if price is not None:
        item["price"] = price
    old_price = _parse_price(row.get("old_price"))
    if old_price is not None:
        item["old_price"] = old_price
        if price is not None and old_price <= price:
            warnings.append("old_price не больше price — скидка не будет показана")

    if row.get("stock") in (None, ""):
        item["stock"] = 0
    else:
        try:
            item["stock"] = max(0, int(float(str(row["stock"]).replace(" ", ""))))
        except (TypeError, ValueError):
            warnings.append(f"некорректный stock {row.get('stock')!r} — записан 0")
            item["stock"] = 0

    if row.get("warranty_months") not in (None, ""):
        try:
            item["warranty_months"] = max(0, int(float(row["warranty_months"])))
        except (TypeError, ValueError):
            warnings.append("некорректный warranty_months — пропущен")

    for f in _BOOL_FIELDS:
        v = _parse_bool(row.get(f), default=None)
        if v is not None:
            item[f] = v
    if "is_active" not in item:
        item["is_active"] = True

    specs = _parse_json_field(row.get("specs"), "dict")
    if specs is None:
        warnings.append("specs не является JSON-объектом — пропущены")
    elif specs:
        item["specs"] = specs
    tags = _parse_json_field(row.get("tags"), "list")
    if tags:
        item["tags"] = tags

    images = [str(row[k]).strip() for k in ("image_1", "image_2", "image_3")
              if row.get(k) and str(row[k]).strip()]
    if images:
        item["images"] = images

    return item, errors, warnings


# ==================== Разбор файлов ====================

def _rows_from_csv(data: bytes) -> list[dict]:
    text = data.decode("utf-8-sig", errors="replace")
    sample = text[:2048]
    delimiter = ";" if sample.count(";") > sample.count(",") else ","
    return list(csv.DictReader(io.StringIO(text), delimiter=delimiter))


def _rows_from_json(data: bytes) -> list[dict]:
    parsed = json.loads(data.decode("utf-8-sig", errors="replace"))
    items = parsed if isinstance(parsed, list) else parsed.get("items") if isinstance(parsed, dict) else None
    if not isinstance(items, list):
        raise ValueError("Ожидается JSON array товаров или {\"items\": [...]}")
    return [r for r in items if isinstance(r, dict)]


def _rows_from_xlsx(data: bytes) -> list[dict]:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    header = next(rows_iter, None)
    if not header:
        return []
    keys = [str(h or "").strip().lower() for h in header]
    out = []
    for values in rows_iter:
        if values is None or all(v in (None, "") for v in values):
            continue
        out.append({k: v for k, v in zip(keys, values) if k})
    wb.close()
    return out


async def _read_rows(file: UploadFile) -> list[dict]:
    data = await file.read()
    if len(data) > MAX_FILE_BYTES:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Файл больше 20 МБ")
    name = (file.filename or "").lower()
    try:
        if name.endswith(".csv") or name.endswith(".txt"):
            rows = _rows_from_csv(data)
        elif name.endswith(".json"):
            rows = _rows_from_json(data)
        elif name.endswith(".xlsx"):
            rows = _rows_from_xlsx(data)
        else:
            raise HTTPException(status.HTTP_400_BAD_REQUEST, "Поддерживаются файлы: .csv, .json, .xlsx")
    except HTTPException:
        raise
    except Exception as e:  # noqa: BLE001 — кривой файл = 400 с пояснением
        raise HTTPException(status.HTTP_400_BAD_REQUEST, f"Не удалось разобрать файл: {str(e)[:200]}")
    if len(rows) > MAX_ROWS:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, f"Не больше {MAX_ROWS} строк за один импорт")
    if not rows:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "В файле нет строк с данными")
    return rows


def _analyze(rows: list[dict], db: Session) -> dict:
    """Общий проход preview/confirm: нормализация, валидация, план действий."""
    # Матчинг регистронезависимый: ключи — lower(sku)
    existing = {
        p.sku.lower(): p
        for p in db.execute(select(Product).where(Product.sku.is_not(None))).scalars()
    }
    plan: list[dict] = []
    report = {"total": len(rows), "created": 0, "updated": 0, "skipped": 0,
              "errors": [], "warnings": [], "items": []}
    seen_skus: set[str] = set()

    for i, raw in enumerate(rows):
        line = i + 2  # человеку удобнее номер строки файла (с учётом заголовка)
        item, errors, warnings = _normalize_row(raw, i)
        if item is None:
            report["skipped"] += 1
            report["errors"].append({"line": line, "sku": raw.get("sku"), "error": "; ".join(errors)})
            continue
        sku = item["sku"]
        if sku.lower() in seen_skus:
            report["skipped"] += 1
            report["errors"].append({"line": line, "sku": sku, "error": "дубль sku в этом же файле"})
            continue
        seen_skus.add(sku.lower())

        product = existing.get(sku.lower())
        action = "update" if product else "create"
        if action == "create":
            if not item.get("title"):
                errors.append("для нового товара нужен title")
            if item.get("price") is None:
                errors.append("для нового товара нужна корректная price")
        if errors:
            report["skipped"] += 1
            report["errors"].append({"line": line, "sku": sku, "error": "; ".join(errors)})
            continue
        for w in warnings:
            report["warnings"].append({"line": line, "sku": sku, "warning": w})

        report[f"{action}d"] += 1
        report["items"].append({
            "line": line, "sku": sku, "action": action,
            "title": item.get("title") or (product.title if product else ""),
            "price": item.get("price") if item.get("price") is not None
                     else (float(product.price) if product else None),
            "stock": item.get("stock", 0),
        })
        plan.append({"action": action, "item": item, "product_id": product.id if product else None})

    report["plan"] = plan
    return report


def _apply_plan(plan: list[dict], db: Session) -> None:
    for step in plan:
        item = dict(step["item"])
        images = item.pop("images", None)
        if step["action"] == "create":
            product = Product(title=item["title"], price=item["price"], source="import")
        else:
            product = db.get(Product, step["product_id"])
            if product is None:
                continue
        for field, value in item.items():
            setattr(product, field, value)
        if images:
            product.images = images
            product.image = images[0]
        product.in_stock = (product.stock or 0) > 0
        if step["action"] == "create":
            db.add(product)
    db.commit()


@router.post("/import/products/preview")
async def import_preview(file: UploadFile = File(...), db: Session = Depends(get_db)):
    rows = await _read_rows(file)
    report = _analyze(rows, db)
    report.pop("plan")           # план — внутренняя кухня, наружу отдаём отчёт
    report["items"] = report["items"][:200]
    return report


@router.post("/import/products/confirm")
async def import_confirm(file: UploadFile = File(...), db: Session = Depends(get_db)):
    rows = await _read_rows(file)
    report = _analyze(rows, db)
    _apply_plan(report.pop("plan"), db)
    report["items"] = report["items"][:200]
    report["applied"] = True
    return report


# ==================== ZIP с фото: авто-матчинг по SKU ====================

_IMG_EXT = {".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".png": "image/png",
            ".webp": "image/webp", ".gif": "image/gif"}


def _sku_from_filename(name: str) -> tuple[str, int]:
    """'IPH15PRO128BLACK-2.jpg' -> ('IPH15PRO128BLACK', 2). '_main'/без суффикса -> 0.

    Порядок: SKU_main/SKU -> 0 (главная), SKU-1 -> 1, SKU-2 -> 2 и т.д.
    """
    stem = PurePosixPath(name.replace("\\", "/")).stem
    lower = stem.lower()
    if lower.endswith("_main") or lower.endswith("-main"):
        return stem[:-5], 0
    for sep in ("-", "_"):
        head, _, tail = stem.rpartition(sep)
        if head and tail.isdigit():
            return head, int(tail)
    return stem, 0


@router.post("/uploads/products/images-zip")
async def upload_images_zip(file: UploadFile = File(...), db: Session = Depends(get_db)):
    """ZIP с фото товаров. Матчинг: имя файла (до расширения) = SKU товара,
    суффиксы -1/-2/_main задают порядок. Первое фото становится главным."""
    data = await file.read()
    if len(data) > MAX_ZIP_BYTES:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "ZIP больше 200 МБ")
    try:
        zf = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Файл не является корректным ZIP-архивом")

    products = {
        (p.sku or "").lower(): p
        for p in db.execute(select(Product).where(Product.sku.is_not(None))).scalars()
    }
    # sku -> [(order, имя файла, content_type, данные)]
    matched: dict[str, list[tuple[int, str, str, bytes]]] = {}
    unmatched: list[str] = []
    errors: list[dict] = []

    for info in zf.infolist():
        if info.is_dir():
            continue
        base = PurePosixPath(info.filename.replace("\\", "/")).name
        if base.startswith(".") or base.startswith("__MACOSX"):
            continue
        ext = ("." + base.rsplit(".", 1)[-1].lower()) if "." in base else ""
        if ext not in _IMG_EXT:
            errors.append({"file": info.filename, "error": f"не изображение ({ext or 'без расширения'})"})
            continue
        if info.file_size > 8 * 1024 * 1024:
            errors.append({"file": info.filename, "error": "файл больше 8 МБ"})
            continue
        sku, order = _sku_from_filename(base)
        product = products.get(sku.lower())
        if product is None:
            unmatched.append(info.filename)
            continue
        try:
            payload = zf.read(info)
        except Exception as e:  # noqa: BLE001
            errors.append({"file": info.filename, "error": str(e)[:100]})
            continue
        matched.setdefault(product.sku, []).append((order, base, _IMG_EXT[ext], payload))

    matched_report = []
    for sku, files in matched.items():
        product = products[sku.lower()]
        files.sort(key=lambda f: f[0])
        urls = [save_image(ctype, payload) for _, _, ctype, payload in files]
        product.images = urls
        product.image = urls[0]
        matched_report.append({"sku": sku, "product_id": product.id,
                               "title": product.title, "images": len(urls)})
    db.commit()

    without_images = [
        {"sku": p.sku, "id": p.id, "title": p.title}
        for p in products.values()
        if not p.images and (not p.image or p.image.startswith("/assets/placeholders/"))
    ]
    return {
        "matched_products": matched_report,
        "unmatched_images": unmatched,
        "products_without_images": without_images[:100],
        "errors": errors,
    }


# ==================== Загрузка одной картинки (баннеры/категории/медиа) ====================

@router.post("/uploads/image", status_code=status.HTTP_201_CREATED)
async def upload_single_image(file: UploadFile = File(...)):
    """Универсальная загрузка картинки: вернёт URL для баннера/категории/товара."""
    from app.core.uploads import MAX_BYTES, is_allowed
    if not is_allowed(file.content_type):
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Только изображения: jpg, png, webp, gif")
    data = await file.read()
    if len(data) > MAX_BYTES:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Файл больше 8 МБ")
    return {"url": save_image(file.content_type, data)}
