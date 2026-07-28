"""Превью прайс-постов без единого обращения к Telegram (v5.6.0).

Источник товаров — БД каталога (как и при боевой генерации) либо XLSX-прайс
через --xlsx. Второй режим существует ради проверки: он позволяет увидеть
посты по файлу ДО того, как каталог импортирован, и сравнить результат с тем,
что получится после импорта.

    python -m app.scripts.preview_price_posts --out /tmp/preview
    python -m app.scripts.preview_price_posts --xlsx price.xlsx --out /tmp/preview

Пишет preview.txt (читаемый вид) и preview.json (структура с клавиатурами).
Ничего не публикует и ничего не меняет.
"""
from __future__ import annotations

import argparse
import json
import sys
from datetime import date
from pathlib import Path

from app.core.config import settings
from app.services.price_posts import (
    SECTIONS,
    catalog_fingerprint,
    navigation_keyboard,
    navigation_text,
    render_all,
    select_products,
)


def load_from_xlsx(path: Path) -> list[dict]:
    import openpyxl

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["products"] if "products" in workbook.sheetnames else workbook.worksheets[0]
    rows = list(sheet.iter_rows(values_only=True))
    header = {name: index for index, name in enumerate(rows[0]) if name}

    def value(row, column, default=None):
        index = header.get(column)
        return row[index] if index is not None and row[index] is not None else default

    products: list[dict] = []
    for row in rows[1:]:
        if not value(row, "sku") and not value(row, "title"):
            continue
        products.append({
            "sku": value(row, "sku"),
            "title": value(row, "title", ""),
            "brand": value(row, "brand"),
            "category": value(row, "category"),
            "subcategory": value(row, "subcategory"),
            "price": value(row, "price", 0),
            "old_price": value(row, "old_price"),
            "is_active": bool(value(row, "is_active", True)),
            "is_hot": bool(value(row, "is_hot", False)),
        })
    return products


def load_from_db() -> list[dict]:
    from app.db.session import SessionLocal
    from app.models.product import Product

    with SessionLocal() as session:
        return [{
            "sku": p.sku, "title": p.title, "brand": p.brand,
            "category": p.category, "subcategory": p.subcategory,
            "price": float(p.price), "old_price": float(p.old_price) if p.old_price else None,
            "is_active": p.is_active, "is_hot": p.is_hot,
        } for p in session.query(Product).all()]


def main() -> int:
    parser = argparse.ArgumentParser(description="Превью прайс-постов канала")
    parser.add_argument("--xlsx", type=Path, help="взять товары из XLSX вместо БД")
    parser.add_argument("--out", type=Path, required=True, help="куда сложить preview.txt/json")
    parser.add_argument("--mini-app-url", default=settings.MINI_APP_URL)
    parser.add_argument("--manager-url", default=settings.MANAGER_RETAIL_URL)
    parser.add_argument("--channel-id", default=settings.TELEGRAM_CHANNEL_ID or "@isellerhub")
    args = parser.parse_args()

    products = load_from_xlsx(args.xlsx) if args.xlsx else load_from_db()
    if not products:
        print("товаров не найдено", file=sys.stderr)
        return 1

    today = date.today()
    posts = render_all(products, today, args.mini_app_url, args.manager_url)
    args.out.mkdir(parents=True, exist_ok=True)

    active = [p for p in products if p.get("is_active", True)]
    covered = {p["sku"] or p["title"] for section in SECTIONS
               for p in select_products(products, section)}
    uncovered = [p for p in active if (p["sku"] or p["title"]) not in covered]

    lines: list[str] = [
        "ПРЕВЬЮ ПРАЙС-ПОСТОВ AI SELLER",
        f"дата: {today:%d.%m.%Y}",
        f"товаров всего: {len(products)}, активных: {len(active)}",
        f"постов: {len(posts)}",
        f"отпечаток каталога: {catalog_fingerprint(products)}",
        "",
    ]
    if uncovered:
        lines += ["!! НЕ ПОПАЛИ НИ В ОДИН РАЗДЕЛ:"]
        lines += [f"   {p['title']} ({p.get('category')} / {p.get('subcategory')})" for p in uncovered]
        lines += [""]

    payload = {
        "generated_at": today.isoformat(),
        "fingerprint": catalog_fingerprint(products),
        "uncovered": [p["title"] for p in uncovered],
        "posts": [],
    }

    for post in posts:
        lines += [
            "=" * 72,
            f"slug: {post.slug}   раздел: {post.section_slug}   товаров: {post.item_count}"
            f"   длина: {len(post.text)}/4096",
            "=" * 72,
            post.text,
            "",
            "кнопки:",
        ]
        for row in post.keyboard:
            lines.append("  [" + "] [".join(
                f"{b['text']} -> {b.get('url') or b['web_app']['url']}" for b in row) + "]")
        lines.append("")
        payload["posts"].append({
            "slug": post.slug, "section_slug": post.section_slug, "title": post.title,
            "text": post.text, "item_count": post.item_count,
            "part": post.part, "parts_total": post.parts_total,
            "length": len(post.text), "keyboard": post.keyboard,
        })

    # Навигация: показываем, как она будет выглядеть, подставив фиктивные
    # message_id — настоящие появятся только после публикации разделов.
    fake_ids = {section.slug: 1000 + index for index, section in enumerate(SECTIONS)}
    nav_keyboard = navigation_keyboard(fake_ids, args.channel_id, args.mini_app_url, args.manager_url)
    lines += ["=" * 72, "НАВИГАЦИОННЫЙ ПОСТ (message_id проставлены условно)", "=" * 72,
              navigation_text(today), "", "кнопки:"]
    for row in nav_keyboard:
        lines.append("  [" + "] [".join(
            f"{b['text']} -> {b.get('url') or b['web_app']['url']}" for b in row) + "]")
    payload["navigation"] = {"text": navigation_text(today), "keyboard": nav_keyboard}

    (args.out / "preview.txt").write_text("\n".join(lines), encoding="utf-8")
    (args.out / "preview.json").write_text(
        json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

    print(f"постов: {len(posts)}, товаров: {len(active)} активных из {len(products)}")
    if uncovered:
        print(f"ВНИМАНИЕ: {len(uncovered)} товаров не попали ни в один раздел")
    for post in posts:
        flag = "  ПРЕВЫШЕН ЛИМИТ" if len(post.text) > 4096 else ""
        print(f"  {post.slug:26} {post.item_count:3} товаров  {len(post.text):5} симв.{flag}")
    print(f"\nзаписано в {args.out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
